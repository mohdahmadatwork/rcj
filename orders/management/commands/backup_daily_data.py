# management/commands/backup_daily_data.py
"""
Django management command to backup all models' data for current date to Excel
Usage: python manage.py backup_daily_data
"""

from django.core.management.base import BaseCommand
from django.apps import apps
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from datetime import datetime, date
import os
from django.conf import settings


class Command(BaseCommand):
    help = 'Creates Excel backup of all models data for current date (24 hours)'

    def add_arguments(self, parser):
        parser.add_argument(
            '--date',
            type=str,
            help='Date in YYYY-MM-DD format (default: today)',
        )
        parser.add_argument(
            '--output',
            type=str,
            help='Output directory path (default: MEDIA_ROOT/backups)',
        )

    def handle(self, *args, **options):
        # Get date parameter or use today
        if options['date']:
            target_date = datetime.strptime(options['date'], '%Y-%m-%d').date()
        else:
            target_date = date.today()

        # Setup output directory
        if options['output']:
            output_dir = options['output']
        else:
            output_dir = os.path.join(settings.MEDIA_ROOT, 'backups')
        
        os.makedirs(output_dir, exist_ok=True)

        # Create workbook
        wb = Workbook()
        wb.remove(wb.active)  # Remove default sheet

        self.stdout.write(self.style.SUCCESS(f'Creating backup for date: {target_date}'))

        # Define models to backup with their date fields
        models_config = [
            {
                'app': 'users',
                'model': 'CustomUser',
                'sheet_name': 'Users',
                'date_field': 'date_joined',
            },
            {
                'app': 'orders',
                'model': 'Order',
                'sheet_name': 'Orders',
                'date_field': 'created_at',
            },
            {
                'app': 'orders',
                'model': 'OrderFile',
                'sheet_name': 'Order Files',
                'date_field': 'uploaded_at',
            },
            {
                'app': 'orders',
                'model': 'OrderLog',
                'sheet_name': 'Order Logs',
                'date_field': 'timestamp',
            },
            {
                'app': 'orders',
                'model': 'Message',
                'sheet_name': 'Messages',
                'date_field': 'created_at',
            },
            {
                'app': 'orders',
                'model': 'Contact',
                'sheet_name': 'Contacts',
                'date_field': 'created_at',
            },
            {
                'app': 'homepage_new',
                'model': 'WorkSample',
                'sheet_name': 'Work Samples',
                'date_field': 'created_at',
            },
            {
                'app': 'homepage_new',
                'model': 'Category',
                'sheet_name': 'Categories',
                'date_field': 'created_at',
            },
            {
                'app': 'news',
                'model': 'NewsItem',
                'sheet_name': 'News Items',
                'date_field': 'created_at',
            },
            {
                'app': 'news',
                'model': 'NewsReadTracker',
                'sheet_name': 'News Read Tracker',
                'date_field': 'read_at',
            },
            {
                'app': 'news',
                'model': 'NewsImage',
                'sheet_name': 'News Images',
                'date_field': 'uploaded_at',
            },
        ]

        total_records = 0

        for config in models_config:
            try:
                # Get model
                Model = apps.get_model(config['app'], config['model'])
                
                # Query data for target date
                filter_kwargs = {f"{config['date_field']}__date": target_date}
                queryset = Model.objects.filter(**filter_kwargs)
                count = queryset.count()

                self.stdout.write(f"Processing {config['sheet_name']}: {count} records")

                # Create sheet
                ws = wb.create_sheet(title=config['sheet_name'])
                
                # Get all field names
                fields = [f.name for f in Model._meta.get_fields() 
                         if not f.many_to_many and not f.one_to_many]

                # Write headers
                self._write_header(ws, fields)

                # Write data
                for row_idx, obj in enumerate(queryset, start=2):
                    for col_idx, field_name in enumerate(fields, start=1):
                        try:
                            value = getattr(obj, field_name)
                            
                            # Handle different field types
                            if hasattr(value, 'pk'):  # ForeignKey
                                value = str(value)
                            elif isinstance(value, (datetime, date)):
                                value = value.isoformat()
                            elif value is None:
                                value = ''
                            else:
                                value = str(value)
                            
                            ws.cell(row=row_idx, column=col_idx, value=value)
                        except Exception as e:
                            ws.cell(row=row_idx, column=col_idx, value='ERROR')
                            self.stdout.write(
                                self.style.WARNING(
                                    f"Error reading {field_name}: {str(e)}"
                                )
                            )

                total_records += count

            except LookupError:
                self.stdout.write(
                    self.style.WARNING(
                        f"Model {config['app']}.{config['model']} not found, skipping..."
                    )
                )
            except Exception as e:
                self.stdout.write(
                    self.style.ERROR(
                        f"Error processing {config['sheet_name']}: {str(e)}"
                    )
                )

        # Save workbook
        filename = f"backup_{target_date.strftime('%Y%m%d')}_{datetime.now().strftime('%H%M%S')}.xlsx"
        filepath = os.path.join(output_dir, filename)
        wb.save(filepath)

        self.stdout.write(
            self.style.SUCCESS(
                f'\n✓ Backup completed successfully!'
                f'\n  Total records: {total_records}'
                f'\n  File saved: {filepath}'
            )
        )

    def _write_header(self, worksheet, headers):
        """Write styled header row"""
        header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        
        for col_idx, header in enumerate(headers, start=1):
            cell = worksheet.cell(row=1, column=col_idx, value=header)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            worksheet.column_dimensions[cell.column_letter].width = 20
